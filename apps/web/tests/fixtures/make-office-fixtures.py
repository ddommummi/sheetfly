#!/usr/bin/env python3
"""XLSX·PPTX 시험용 문서를 만든다(openpyxl, python-pptx; 둘 다 MIT/BSD). 결과 파일은 저장소에 함께 둔다 — 다시 만들려면:
   python3 make-office-fixtures.py   (같은 폴더에 sample.xlsx, sample.pptx를 쓴다. sample.xls/ppt는 LibreOffice로 변환:
   soffice --headless --convert-to xls sample.xlsx ; soffice --headless --convert-to ppt sample.pptx)"""
import datetime, struct, zlib, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

here = os.path.dirname(os.path.abspath(__file__))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = '매출'
ws.append(['항목', '1월', '2월', '합계'])
for i, n in enumerate(['사과', '배', '포도']):
    ws.append([n, 100 + i * 10, 200 + i * 5, 300 + i * 15])
ws['A1'].font = Font(bold=True, color='FFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='4472C4')
ws.merge_cells('A6:D6'); ws['A6'] = '병합된 셀 입니다'; ws['A6'].alignment = Alignment(horizontal='center')
ws['B2'].number_format = '#,##0.00'; ws['C3'] = 0.256; ws['C3'].number_format = '0.0%'
ws2 = wb.create_sheet('넓은 시트')
ws2.append([f'열{c}' for c in range(1, 41)])            # 40열: 쪽 너비를 넘는다
for r in range(2, 160):                                  # 159행: 여러 쪽
    ws2.append([f'행{r}-' + ('가' * (c % 5)) if c % 7 == 0 else r * c for c in range(1, 41)])
ws3 = wb.create_sheet('Third'); ws3['A1'] = 'third sheet'
wb.save(os.path.join(here, 'sample.xlsx'))

def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def ch(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + ch(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + ch(b'IDAT', zlib.compress(raw)) + ch(b'IEND', b'')
open(os.path.join(here, '_red.png'), 'wb').write(png(40, 30, (200, 30, 30)))

p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = '제목 슬라이드'; s.placeholders[1].text = '부제목 Subtitle'
s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = '목록'
tf = s.placeholders[1].text_frame; tf.text = '첫째 항목'; tf.add_paragraph().text = '둘째 항목'
s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = '표와 그림'
t = s.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
for r in range(3):
    for c in range(3): t.cell(r, c).text = f'{r},{c} 값'
s.shapes.add_picture(os.path.join(here, '_red.png'), Inches(8), Inches(2), Inches(2), Inches(1.5))
sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(3), Inches(1)); sh.text_frame.text = '도형 텍스트'
# 4번째: 제목 자리 표시자가 직접 위치를 가진다(그림 아래) — 레이아웃 기본 위치(위쪽)와 다르다. 실제 발표 자료에서 흔한 모양.
s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = 'AI 에이전트'
s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(5), Inches(6), Inches(3), Inches(1)
s.shapes.add_picture(os.path.join(here, '_red.png'), Inches(4), Inches(0.5), Inches(5), Inches(5))
p.save(os.path.join(here, 'sample.pptx'))
os.remove(os.path.join(here, '_red.png'))
