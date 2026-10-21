# 10만 줄 통합 문서: 값 열·수식(곱셈·IF)·VLOOKUP(정확히 일치, 1,000줄 표)·SUMIF/AVERAGE/COUNTIF 요약
import openpyxl
wb = openpyxl.Workbook(write_only=True)
ws = wb.create_sheet('Data'); N = 100000
ws.append(['값', 'x1.1', '구분', '합계', '조회'])
for r in range(2, N + 2):
    ws.append([r * 7 % 1000, f'=A{r}*1.1', f'=IF(B{r}>500,"high","low")', None, f'=VLOOKUP(A{r},Tbl!$A$2:$B$1001,2,FALSE)'])
t = wb.create_sheet('Tbl'); t.append(['k', 'v'])
for i in range(1, 1001): t.append([i - 1, f'=A{i+1}*3'])
s = wb.create_sheet('Sum')
s.append(['sumhigh', f'=SUMIF(Data!C2:C{N+1},"high",Data!B2:B{N+1})']); s.append(['avg', f'=AVERAGE(Data!B2:B{N+1})']); s.append(['cnt', f'=COUNTIF(Data!C2:C{N+1},"low")'])
wb.save('big.xlsx')
