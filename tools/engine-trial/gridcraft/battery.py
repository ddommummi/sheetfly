import openpyxl, sys, datetime
wb = openpyxl.Workbook()
d = wb.active; d.title = 'D'
# 데이터 블록(한국어 포함)
names = ['김철수','이영희','박민수','최지은','정우성','한소희','오세훈','신민아']
depts = ['영업','개발','영업','인사','개발','영업','인사','개발']
amts  = [1200000, 850000, 2300000, 640000, 1750000, 980000, 310000, 2100000]
d.append(['이름','부서','금액','입사일','비고'])
for i,(n,p,a) in enumerate(zip(names,depts,amts)):
    d.append([n,p,a,datetime.date(2020+i%4, 1+i, 3+i),'우수' if a>1500000 else '보통'])
for i in range(10, 40): d.cell(row=i, column=7, value=i*1.5)        # G10:G39 숫자
d['I1']=' Hello  World '; d['I2']='Excel,Word,PPT'; d['I3']='가나다라마바사'; d['I4']='2024-03-15'; d['I5']='1,234.5'; d['I6']='AbC dEf'
d['J1']=-12.345; d['J2']=1234567.891; d['J3']=0.075; d['J4']=45000; d['J5']=3.14159265358979
d['K1']=True; d['K2']=False; d['K3']=None
cases = []
def add(name, f): cases.append((name, f))
# 수학
for fn in ['ABS','SQRT','INT','ROUND','ROUNDUP','ROUNDDOWN','CEILING','FLOOR','TRUNC','SIGN','EXP','LN','LOG10','FACT','SIN','COS','TAN','ATAN','DEGREES','RADIANS','EVEN','ODD']:
    arg = {'ROUND':'D!J2,-2','ROUNDUP':'D!J1,1','ROUNDDOWN':'D!J1,1','CEILING':'D!J1,0.5','FLOOR':'D!J1,0.5','SQRT':'D!G12','LN':'D!G12','LOG10':'D!G12','FACT':'5','ABS':'D!J1','INT':'D!J1','TRUNC':'D!J1,1','SIGN':'D!J1','EXP':'1','SIN':'D!J3','COS':'D!J3','TAN':'D!J3','ATAN':'D!J3','DEGREES':'D!J5','RADIANS':'180','EVEN':'D!J5','ODD':'D!J5'}[fn]
    add(fn, f'={fn}({arg})')
add('MOD','=MOD(-7,3)'); add('POWER','=POWER(2,10)'); add('PRODUCT','=PRODUCT(D!G10:G14)'); add('SUMSQ','=SUMSQ(D!G10:G14)'); add('GCD','=GCD(24,36)'); add('LCM','=LCM(4,6)'); add('QUOTIENT','=QUOTIENT(17,5)'); add('MROUND','=MROUND(17,5)'); add('COMBIN','=COMBIN(8,3)'); add('PERMUT','=PERMUT(8,3)'); add('PI','=PI()'); add('SUMPRODUCT','=SUMPRODUCT(D!G10:G14,D!G15:G19)')
# 통계/집계
for fn in ['SUM','AVERAGE','MIN','MAX','COUNT','COUNTA','MEDIAN','STDEV','STDEV.P','VAR','VAR.S','LARGE','SMALL','MODE','COUNTBLANK']:
    rng='D!C2:C9'
    extra = ',2' if fn in ('LARGE','SMALL') else ''
    if fn=='COUNTBLANK': rng='D!K1:K3'
    add(fn, f'={fn}({rng}{extra})')
add('RANK','=RANK(D!C3,D!C2:C9)'); add('PERCENTILE','=PERCENTILE(D!C2:C9,0.3)'); add('QUARTILE','=QUARTILE(D!C2:C9,1)'); add('CORREL','=CORREL(D!G10:G20,D!G11:G21)'); add('SLOPE','=SLOPE(D!G10:G20,D!G11:G21)'); add('FORECAST','=FORECAST(50,D!G10:G20,D!G11:G21)'); add('AVERAGEIF','=AVERAGEIF(D!B2:B9,"영업",D!C2:C9)'); add('AVERAGEIFS','=AVERAGEIFS(D!C2:C9,D!B2:B9,"개발",D!C2:C9,">1000000")')
add('SUMIF','=SUMIF(D!B2:B9,"영업",D!C2:C9)'); add('SUMIFS','=SUMIFS(D!C2:C9,D!B2:B9,"영업",D!C2:C9,">1000000")'); add('COUNTIF','=COUNTIF(D!B2:B9,"영업")'); add('COUNTIFS','=COUNTIFS(D!B2:B9,"개발",D!C2:C9,">=1750000")'); add('MAXIFS','=MAXIFS(D!C2:C9,D!B2:B9,"인사")'); add('MINIFS','=MINIFS(D!C2:C9,D!B2:B9,"영업")')
add('SUBTOTAL','=SUBTOTAL(9,D!C2:C9)'); add('AGGREGATE','=AGGREGATE(9,6,D!C2:C9)')
# 논리
add('IF','=IF(D!C2>1000000,"고액","소액")'); add('IFS','=IFS(D!C2>2000000,"A",D!C2>1000000,"B",TRUE,"C")'); add('AND','=AND(D!C2>1,D!C3>1)'); add('OR','=OR(D!C2<1,D!C3>1)'); add('NOT','=NOT(D!K1)'); add('XOR','=XOR(TRUE,FALSE)'); add('IFERROR','=IFERROR(1/0,"오류")'); add('IFNA','=IFNA(NA(),"없음")'); add('SWITCH','=SWITCH(D!B2,"영업",1,"개발",2,0)'); add('ISNUMBER','=ISNUMBER(D!C2)'); add('ISTEXT','=ISTEXT(D!A2)'); add('ISBLANK','=ISBLANK(D!K3)'); add('ISERROR','=ISERROR(1/0)'); add('ISEVEN','=ISEVEN(4)'); add('ISODD','=ISODD(4)')
# 텍스트
for fn,a in [('LEN','D!I3'),('LEFT','D!I3,2'),('RIGHT','D!I3,2'),('MID','D!I3,2,3'),('UPPER','D!I6'),('LOWER','D!I6'),('PROPER','D!I6'),('TRIM','D!I1'),('CLEAN','D!I1'),('VALUE','D!I5'),('FIND','"다",D!I3'),('SEARCH','"D",D!I6'),('SUBSTITUTE','D!I2,",","/"'),('REPLACE','D!I3,2,2,"**"'),('REPT','"가",3'),('CONCATENATE','D!A2,"-",D!B2'),('CONCAT','D!A2,D!B2'),('TEXTJOIN','",",TRUE,D!A2:A4'),('EXACT','D!I6,"AbC dEf"'),('CODE','"A"'),('CHAR','65'),('UNICODE','D!I3'),('UNICHAR','44032'),('LENB','D!I3'),('LEFTB','D!I3,4'),('T','D!I3'),('N','D!C2'),('FIXED','D!J2,1'),('DOLLAR','D!J2,0'),('TEXTBEFORE','D!I2,","'),('TEXTAFTER','D!I2,","')]:
    add(fn, f'={fn}({a})')
for i,fmt in enumerate(['#,##0','#,##0.00','0.0%','yyyy-mm-dd','yyyy"년" m"월" d"일"','aaa','aaaa','#,##0"원"','[>=1000000]#,##0,,"백만";0','hh:mm','0.00E+00','@']):
    v = 'D!J2' if i not in (3,4,5,6,10-6) else 'D!D2'
    if fmt in ('yyyy-mm-dd','yyyy"년" m"월" d"일"','aaa','aaaa'): v='D!D2'
    if fmt=='hh:mm': v='0.75'
    if fmt=='0.0%': v='D!J3'
    if fmt=='@': v='D!J1'
    add(f'TEXT[{fmt}]', f'=TEXT({v},"{fmt.replace(chr(34), chr(34)*2)}")')
# 날짜
add('TODAY형(고정)','=DATE(2024,2,29)'); add('YEAR','=YEAR(D!D2)'); add('MONTH','=MONTH(D!D3)'); add('DAY','=DAY(D!D4)'); add('WEEKDAY','=WEEKDAY(D!D2,2)'); add('EDATE','=EDATE(D!D2,13)'); add('EOMONTH','=EOMONTH(D!D2,1)'); add('DATEDIF','=DATEDIF(D!D2,D!D9,"M")'); add('DAYS','=DAYS(D!D9,D!D2)'); add('NETWORKDAYS','=NETWORKDAYS(D!D2,D!D3)'); add('WORKDAY','=WORKDAY(D!D2,10)'); add('YEARFRAC','=YEARFRAC(D!D2,D!D9)'); add('WEEKNUM','=WEEKNUM(D!D4)'); add('TIME','=TIME(13,45,30)'); add('HOUR','=HOUR(0.75)'); add('DATEVALUE','=DATEVALUE(D!I4)'); add('TIMEVALUE','=TIMEVALUE("18:30")'); add('ISOWEEKNUM','=ISOWEEKNUM(D!D4)')
# 조회
add('VLOOKUP','=VLOOKUP("박민수",D!A2:E9,3,FALSE)'); add('HLOOKUP','=HLOOKUP("금액",D!A1:E9,3,FALSE)'); add('LOOKUP','=LOOKUP(2000000,D!C2:C9)'); add('INDEX','=INDEX(D!A2:E9,3,2)'); add('MATCH','=MATCH("한소희",D!A2:A9,0)'); add('MATCH근사','=MATCH(1000000,D!G10:G39,1)'); add('XLOOKUP','=XLOOKUP("오세훈",D!A2:A9,D!C2:C9,"없음")'); add('XMATCH','=XMATCH("신민아",D!A2:A9)'); add('CHOOSE','=CHOOSE(2,"a","b","c")'); add('OFFSET','=OFFSET(D!A1,2,2)'); add('INDIRECT','=INDIRECT("D!C3")'); add('ROW','=ROW(D!C5)'); add('COLUMN','=COLUMN(D!C5)'); add('ROWS','=ROWS(D!A2:E9)'); add('COLUMNS','=COLUMNS(D!A2:E9)'); add('ADDRESS','=ADDRESS(3,4)'); add('TRANSPOSE합','=SUM(TRANSPOSE(D!C2:C4))'); add('FILTER','=SUM(FILTER(D!C2:C9,D!B2:B9="영업"))'); add('SORT','=INDEX(SORT(D!C2:C9),1)'); add('UNIQUE','=ROWS(UNIQUE(D!B2:B9))'); add('SEQUENCE','=SUM(SEQUENCE(10))'); add('LET','=LET(x,5,y,x*2,x+y)')
# 재무
add('PMT','=PMT(0.05/12,36,-10000000)'); add('FV','=FV(0.03,10,-100000)'); add('PV','=PV(0.04,5,-1000)'); add('NPV','=NPV(0.1,-1000,300,400,500,600)'); add('IRR','=IRR(D!G10:G15)'); add('RATE','=RATE(48,-200,8000)'); add('NPER','=NPER(0.05,-100,1000)'); add('IPMT','=IPMT(0.05/12,1,36,-10000000)'); add('PPMT','=PPMT(0.05/12,1,36,-10000000)'); add('SLN','=SLN(10000,1000,5)'); add('DB','=DB(10000,1000,5,1)'); add('SYD','=SYD(10000,1000,5,2)')
# 연산자/기타
add('연결','=D!A2&"("&D!B2&")"'); add('거듭','=2^0.5'); add('퍼센트','=50%*3'); add('비교','=D!C2>D!C3'); add('배열곱','=SUMPRODUCT((D!B2:B9="영업")*(D!C2:C9))'); add('중첩IF','=IF(D!C2>2000000,"S",IF(D!C2>1000000,"A","B"))'); add('범위이름식','=SUM(D!C2:C9)/COUNT(D!C2:C9)'); add('오류전파','=1/0+1'); add('텍스트수','="12"+3'); add('빈셀','=D!K3+1'); add('불리언수','=TRUE+1')
dd = wb.create_sheet('T'); dd.append(['이름','수식'])
XL = ['IFS','XOR','IFNA','SWITCH','CONCAT','TEXTJOIN','UNICODE','UNICHAR','DAYS','ISOWEEKNUM','XLOOKUP','XMATCH','LET','MAXIFS','MINIFS','STDEV.P','VAR.S','TEXTBEFORE','TEXTAFTER','SEQUENCE','AGGREGATE']
XLWS = ['FILTER','SORT','UNIQUE']
import re
def fix(f):
    for n in XLWS: f = re.sub(r'(?<![\w.])'+n+r'\(', '_xlfn._xlws.'+n+'(', f)
    for n in XL: f = re.sub(r'(?<![\w.])'+re.escape(n)+r'\(', '_xlfn.'+n+'(', f)
    return f
for i,(n,f) in enumerate(cases, start=2):
    f = fix(f)
    dd.cell(row=i, column=1, value=n); dd.cell(row=i, column=2, value=f)
wb.save(sys.argv[1]); print(len(cases), 'cases')
